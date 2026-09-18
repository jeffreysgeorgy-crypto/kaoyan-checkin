# -*- coding: utf-8 -*-
"""
gen_outline.py —— 生成《基于 openJiuwen 的个人学习规划 Agent》解决方案大纲（.docx）

用法：  py gen_outline.py
输出：  D:\\XX\\解决方案大纲_基于openJiuwen的个人学习规划Agent.docx
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"D:\XX\解决方案大纲_基于openJiuwen的个人学习规划Agent.docx"

BODY = "宋体"
HEAD = "黑体"
GRAY = RGBColor(0x80, 0x80, 0x80)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x00, 0x00, 0x00)


def _font(run, east=BODY, bold=False, size=12, color=None, italic=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), east)


def h(doc, text, level=1):
    sizes = {0: 22, 1: 16, 2: 14, 3: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    run = p.add_run(text)
    _font(run, east=HEAD, bold=True, size=sizes.get(level, 12), color=DARK if level <= 2 else BLUE)
    return p


def p(doc, text, indent=0, bold=False, italic=False, color=None, size=12):
    para = doc.add_paragraph()
    if indent:
        para.paragraph_format.left_indent = Cm(0.5 * indent)
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.3
    run = para.add_run(text)
    _font(run, east=BODY, bold=bold, italic=italic, color=color, size=size)
    return para


def bullet(doc, text, level=0):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(0.75 + 0.5 * level)
    para.paragraph_format.first_line_indent = Cm(-0.4)
    para.paragraph_format.space_after = Pt(2)
    para.paragraph_format.line_spacing = 1.25
    run = para.add_run("• " + text)
    _font(run, east=BODY, size=12)
    return para


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(htext)
        _font(run, east=HEAD, bold=True, size=11)
        # 表头底色
        tcPr = hdr[i]._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "DEEAF6")
        tcPr.append(shd)
    for r in rows:
        cells = t.add_row().cells
        for i, ctext in enumerate(r):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(ctext)
            _font(run, east=BODY, size=10.5)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def need(doc, *items):
    """章节末尾：需要补充的材料。"""
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(6)
    run = para.add_run("【需要补充的材料】")
    _font(run, east=HEAD, bold=True, size=11, color=BLUE)
    for it in items:
        b = doc.add_paragraph()
        b.paragraph_format.left_indent = Cm(0.5)
        b.paragraph_format.space_after = Pt(2)
        r = b.add_run("· " + it)
        _font(r, east=BODY, size=11, color=GRAY)


doc = Document()
# 正文默认样式
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY)

# ============================================================
# 封面
# ============================================================
for _ in range(4):
    doc.add_paragraph()
cp = doc.add_paragraph()
cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cp.add_run("基于 openJiuwen 的个人学习规划 Agent")
_font(r, east=HEAD, bold=True, size=26)
cp2 = doc.add_paragraph()
cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cp2.add_run("—— 解决方案 ——")
_font(r, east=HEAD, bold=True, size=18)
for _ in range(6):
    doc.add_paragraph()

cover = [
    ("项目名称", "基于 openJiuwen 的个人学习规划 Agent"),
    ("赛　　题", "基于 openJiuwen 的个人学习规划 Agent"),
    ("赛　　道", "产业命题赛道（企业命题组）"),
    ("团队名称", "【待补充】"),
    ("所属学校", "【待补充】"),
    ("团队成员", "【待补充】"),
    ("指导教师", "【待补充】"),
    ("申报日期", "2026 年【待补充】月"),
]
for k, v in cover:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after = Pt(8)
    r = para.add_run(k + "：")
    _font(r, east=HEAD, bold=True, size=14)
    r2 = para.add_run(v)
    _font(r2, east=BODY, size=14)
doc.add_page_break()

# ============================================================
# 第一章 执行摘要
# ============================================================
h(doc, "第一章　执行摘要", 1)
p(doc, "本方案面向 2027 考研备考场景（数学一 + 408 计算机统考），基于开源 Agent 框架 openJiuwen，构建一个能「记住用户学习状态、动态规划学习路径、根据反馈调整计划」的个人学习规划 Agent。核心价值是让 AI 学习规划从「一次性生成静态课表」走向「持续动态、可解释、可信赖」的闭环。")

h(doc, "1.1 命题理解与解题思路", 2)
h(doc, "1.1.1 所选命题概述", 3)
p(doc, "本方案选择「基于 openJiuwen 的个人学习规划 Agent」命题。openJiuwen 是昇腾开源 Agent 框架（Apache-2.0，新 API：AgentCard + ReActAgentConfig + ReActAgent），命题要求参赛者在真实场景中，用该框架完成一个能记忆、能规划、能动态调整的 Agent 应用。")
h(doc, "1.1.2 命题理解深度（命题背后的真实痛点）", 3)
p(doc, "命题表面是「用框架做一个学习规划 Agent」，其背后的真实痛点是当前学习工具的三大共性缺陷：")
bullet(doc, "静态：计划一经生成即固化，不随执行情况（完成/未完成）变化，用户偏离计划后只能靠人工重排；")
bullet(doc, "无记忆：不记住用户的历史完成率、连续失败、薄弱知识点与反思情绪，每次规划都从零开始；")
bullet(doc, "不可解释：AI 给出的调整是「黑盒」，用户不知道「为什么这么调、依据是什么」，难以建立信任。")
p(doc, "因此，命题的深层意图是考察「Agent 编排 + 领域规则」的工程化落地能力，而非单纯的模型调用。")
h(doc, "1.1.3 解题总体思路（三层架构：规则 + LLM + 调度）", 3)
p(doc, "采用「感知层 → 记忆层 → 决策层」三层解耦架构：")
bullet(doc, "感知层：HTML 打卡系统（打卡 / 错题 / 专注 / 单词 / 反思），是学习事实的来源；")
bullet(doc, "记忆层：memory.json 学习记忆快照（画像 / 计划 / 各科掌握度 / 连续失败 / 反思原文），随打卡动态更新；")
bullet(doc, "决策层：openJiuwen ReActAgent 调度 8 个确定性 Skill（skills.py），输出可解释的决策与调整。")
p(doc, "核心原则是「规则负责算对，LLM 负责表达」：8 个 Skill 全部由确定性规则实现，可复现、可审计；LLM 只在 openJiuwen 的 ReAct 循环中按需调度 Skill 并生成决策日志，从根上避免幻觉。")

h(doc, "1.2 解决方案概述", 2)
h(doc, "1.2.1 核心方案内容（8 个 Skill）", 3)
p(doc, "系统提供 8 个可解释 Skill，覆盖学习规划的全生命周期：diagnose（学习诊断）、replan（动态重规划）、allocate（多科目资源再分配）、interpret_feedback（反思理解）、proactive_scan（未来 7 天负荷扫描）、reschedule_for_calendar_change（课表变动重排）、decompose_goal（目标拆解）、aggregate_resources（资源聚合）。详见第三章 3.2.2。")
h(doc, "1.2.2 方案先进性（可解释、容错、多级兜底）", 3)
bullet(doc, "可解释：每次调整都带 reason（为什么）与结构化的 evidence（连续失败天数 / 完成率 / 反思原文 / 结论），可审计、可回放；")
bullet(doc, "容错降级：LLM 不可用（无 key / 超时 / 连续失败）时自动降级为关键词匹配，任何环境都能跑通完整流程；")
bullet(doc, "多级兜底：干预随「连续失败天数」与「重规划次数」逐级加重，重规划 ≥2 次后换策略并给出三个方向，最后「把人请回来」，避免自动重规划死循环。")

h(doc, "1.3 匹配度与可行性", 2)
h(doc, "1.3.1 方案与命题匹配度（6 答题点对照表）", 3)
table(doc,
      ["答题点", "方案实现位置"],
      [
          ["① 用户画像与学习记忆存储更新", "memory.json 记忆快照 + 打卡闭环更新（见 3.2.1、7.2）"],
          ["② 至少 2 个 Skill（输入/输出/调用条件）", "skills.py 提供 8 个 Skill（见 3.2.2）"],
          ["③ 动态调整案例（连续 3 天未完成）", "demo.py 逐日演示（见 3.3.2）"],
          ["④ 学习计划生成前后对比", "adjustments[] 的 before/after/reason/evidence（见 3.3.2）"],
          ["⑤ openJiuwen 承担的作用", "AgentCard + ReActAgent + @tool 编排（见 3.2.1）"],
          ["⑥ 个性化与可解释性", "掌握度公式 + reason/evidence 证据链（见 3.2.3）"],
      ],
      widths=[5.5, 9.5])
h(doc, "1.3.2 方案可行性（技术 / 经济 / 时间）", 3)
bullet(doc, "技术可行：全部采用开源成熟栈（openJiuwen / FastAPI / HTML+JS），无自研底层依赖；")
bullet(doc, "经济可行：离线可跑（确定性规则无需 API），LLM 可选（DeepSeek / 硅基流动按量付费），无额外硬件成本；")
bullet(doc, "时间可行：系统已完成开发，进入答辩打磨阶段（见 5.2 里程碑）。")

h(doc, "1.4 团队简介", 2)
p(doc, "【待补充：团队成员构成、专业背景、分工与既往成果。】")
h(doc, "1.5 预期成效", 2)
bullet(doc, "完成赛题 6 个答题点全覆盖的可演示系统；")
bullet(doc, "形成「记忆 + 动态调整 + 可解释」的完整产品闭环；")
bullet(doc, "可作为考研学习工具的 MVP，具备产品化潜力。")
need(doc, "团队成员信息与分工", "指导教师信息", "如有既往获奖 / 项目经历，一并补充")

# ============================================================
# 第二章 命题分析与产业认知
# ============================================================
h(doc, "第二章　命题分析与产业认知", 1)
h(doc, "2.1 产业认知（在线教育 / 智能学习工具行业现状）", 2)
p(doc, "AI 与教育进入落地期，但多数「AI 学习规划」仍停留在两个层次：一是「一次性生成课表」（静态，不随执行变化）；二是「通用问答」（缺乏对个人学习历史的记忆与个性化）。真正具备「记忆 + 动态调整 + 可解释」闭环的产品仍是空白，这正是本命题的价值空间。")
p(doc, "（注：在线教育 / 智能学习工具的行业规模、增长率等具体数据，为避免虚构，此处标注待补充，将在正式稿中引用权威来源。）")
h(doc, "2.2 命题企业分析（openJiuwen 的定位与命题意图）", 2)
p(doc, "openJiuwen 是昇腾开源 Agent 框架（Apache-2.0），提供 AgentCard、ReActAgentConfig、ReActAgent、@tool 等新 API，定位是「Agent 决策编排层」。命题意图是考察参赛者能否将「框架能力」与「领域规则」结合，做出有真实价值、可解释、可落地的 Agent 应用，而非堆砌模型调用。")
h(doc, "2.3 命题需求深度剖析", 2)
h(doc, "2.3.1 命题表面需求（赛题原文 + 逐条对齐）", 3)
p(doc, "【命题背景】", bold=True)
p(doc, "本科生在学习编程、数学、英语、科研入门等任务时，常常面临目标不清、计划难坚持、资源分散、学习进度无法动态调整等问题。传统学习计划通常是静态表格，一旦用户进度落后、掌握情况变化或目标调整，就很难自动更新。因此需要构建一个能够记住用户学习状态、动态规划学习路径、根据反馈调整计划的个人学习规划 Agent。")
p(doc, "【命题内容】", bold=True)
p(doc, "基于 openJiuwen 框架，设计并实现一个个人学习规划 Agent，用于帮助学生完成一个阶段性学习目标，例如：30 天入门 Python、2 周准备数学建模竞赛、1 个月完成机器学习基础学习。系统需要具备：用户画像与学习记忆。")
p(doc, "【答题要求】逐条对齐本方案实现：", bold=True)
table(doc,
      ["答题要求（赛题原文）", "本方案实现"],
      [
          ["必须展示用户学习记忆如何存储与更新", "memory.json 记忆快照 + update_memory_after_checkin 打卡闭环更新"],
          ["必须设计至少 2 个 Skill，并说明输入、输出和调用条件", "skills.py 提供 8 个 Skill，详见 3.2.2"],
          ["必须包含一个动态调整案例（例如：用户连续 3 天未完成任务，Agent 如何重规划）", "demo.py Day 3：连续 3 天未完成 → 减少时长 + 拆分任务，详见 3.3.2"],
          ["必须展示学习计划生成前后的对比", "adjustments[] 的 before / after / reason / evidence，详见 3.3.2"],
          ["必须说明 openJiuwen 在其中承担的作用", "AgentCard + ReActAgent + @tool 决策编排层，详见 3.2.1"],
          ["输出需体现个性化与可解释性，例如为什么安排这个任务、为什么调整计划、依据了哪些历史记录等", "掌握度公式 + reason/evidence 证据链，详见 3.2.3"],
      ],
      widths=[8, 7])
h(doc, "2.3.2 命题深层痛点", 3)
bullet(doc, "静态课表：计划不随执行反馈调整；")
bullet(doc, "无记忆：不记录历史完成、薄弱点与情绪；")
bullet(doc, "不能动态调整：缺乏「诊断 → 重规划 → 验证」闭环；")
bullet(doc, "不可解释：调整缺乏依据，用户不信任。")
h(doc, "2.3.3 需求优先级排序", 3)
table(doc,
      ["优先级", "需求", "理由"],
      [
          ["P0", "学习记忆存储与更新", "一切动态调整的前提"],
          ["P1", "动态重规划（诊断 + 调整）", "命题核心「动态」"],
          ["P1", "可解释性（reason/evidence）", "命题核心「可解释」"],
          ["P2", "多 Skill 扩展（分配/扫描/重排）", "体现框架编排能力"],
          ["P2", "个性化（掌握度公式/画像）", "从通用到个体"],
      ],
      widths=[2, 6, 7])
h(doc, "2.4 资源需求分析", 2)
bullet(doc, "开发资源：开发已完成，主要在答辩材料打磨；")
bullet(doc, "数据资源：用户打卡 / 反思数据（demo 用预置 + 运行时注入两类，source 字段可审计）；")
bullet(doc, "算力资源：LLM 调用可选（按量付费），离线降级无需算力。")
need(doc, "行业权威数据来源（市场规模 / 增长率，标注出处）")

# ============================================================
# 第三章 解决方案（核心）
# ============================================================
h(doc, "第三章　解决方案", 1)
p(doc, "本章为方案核心，重点说明「解题理念、技术方案、8 个 Skill 的输入输出与调用条件、实施方案、需求匹配度」五个方面。")

h(doc, "3.1 解题理念与创新思路", 2)
h(doc, "3.1.1 解题理念（让系统记住用户，让调整有据可查）", 3)
p(doc, "一句话概括：让系统「记住用户」，让每一次调整「有据可查」。记忆（memory）是动态调整的前提，证据（evidence）是可解释的保证。二者共同构成与「静态课表、黑盒 AI」的本质区别。")
h(doc, "3.1.2 创新思路（三层解耦、多级兜底、证据驱动）", 3)
bullet(doc, "三层解耦：感知 / 记忆 / 决策分离，任一层可独立替换为生产级实现；")
bullet(doc, "多级兜底：干预逐级加重（标红 → 拆分减时长 → 降级难度 → 换策略 + 三方向 → 人工确认），杜绝死循环；")
bullet(doc, "证据驱动：每条调整绑定结构化 evidence（连续失败天数 / 完成率 / 反思原文 / 结论），杜绝「模型瞎编理由」。")
h(doc, "3.1.3 创新先进性（对比传统工具）", 3)
table(doc,
      ["维度", "传统学习工具", "本方案"],
      [
          ["计划形态", "静态课表，生成即固化", "动态调整，随执行反馈重排"],
          ["记忆", "无 / 仅存打卡记录", "结构化学习记忆（掌握度/连续失败/反思）"],
          ["可解释性", "黑盒，无依据", "reason + evidence，可审计"],
          ["容错", "依赖在线模型", "LLM 不可用自动降级规则引擎"],
          ["兜底", "一次调整 / 反复同一招", "多级兜底 + 换策略 + 人工确认"],
      ],
      widths=[3, 6, 6])

h(doc, "3.2 技术方案", 2)
h(doc, "3.2.1 技术路线图", 3)
p(doc, "核心数据流（一次完整闭环）：")
table(doc,
      ["步骤", "环节", "说明"],
      [
          ["1", "用户打卡", "HTML 前端记录打卡 / 错题 / 专注 / 单词 / 反思"],
          ["2", "记忆更新", "update_memory_after_checkin 重算掌握度 / 连续失败 / 完成率"],
          ["3", "决策调度", "openJiuwen ReActAgent 思考 → 调用 Skill → 观察"],
          ["4", "Skill 执行", "skills.py 确定性规则算出诊断 / 调整 / 分配 / 扫描 / 重排"],
          ["5", "生成新计划", "new_plan + reason + evidence 写回前端，今日时间轴更新"],
      ],
      widths=[1.5, 4, 9.5])
p(doc, "技术栈：Python + FastAPI（后端，9 个 POST 接口）+ openJiuwen（决策编排）+ HTML/JS（前端，localStorage 持久化）+ Supabase（数据，规划中）。")
h(doc, "3.2.2 核心技术 / 方法（8 个 Skill 详述）", 3)
p(doc, "8 个 Skill 全部为确定性规则，输入输出与调用条件如下（每个 Skill 单独一节）。其中前 6 个已接入 HTML 前端入口按钮，decompose_goal / aggregate_resources 两个当前由 demo.py 演示，属「引擎已具备、UI 待补」的边界。")

h(doc, "3.2.2.1 Skill 1：diagnose（学习诊断）", 4)
table(doc,
      ["项", "内容"],
      [
          ["输入", "learning_memory（各科目掌握度 / 连续失败天数 / 近 7 天完成率 / 弱知识点 / 反思原文）"],
          ["输出", "诊断报告：overall_status + alert_subjects[]（severity / 原因 / 主因 / 反思引用）+ healthy_subjects[]"],
          ["调用条件", "① 生成新计划前；② 某科目连续失败达到阈值、触发重规划前"],
          ["关键逻辑", "严重度判定（连续失败≥5 或≥3 天=high；≥2 天或完成率<50%=medium；其余=low）；多维度归因（方法不当/情绪干扰/精力不足/目标不清晰/时间投入不足/任务过载/基础薄弱/环境干扰）"],
      ],
      widths=[3, 12])

h(doc, "3.2.2.2 Skill 2：replan（动态重规划）", 4)
table(doc,
      ["项", "内容"],
      [
          ["输入", "current_plan + diagnosis + learning_memory + replan_count（已重规划次数）"],
          ["输出", "(new_plan, adjustments)；每条 adjustment 含 task_id / subject / level / before / after / reason / evidence"],
          ["调用条件", "诊断出预警科目后，对计划做针对性调整"],
          ["关键逻辑", "分级干预：medium 仅标红提醒；high(连续3天) 减时长+拆分；high(连续≥5天) 降级难度；已重规划≥2次 换任务类型+给出三个方向（推荐 A）"],
      ],
      widths=[3, 12])

h(doc, "3.2.2.3 Skill 3：allocate（多科目资源再分配）", 4)
table(doc,
      ["项", "内容"],
      [
          ["输入", "learning_memory + current_plan + daily_available_hours（每日可用时长）"],
          ["输出", "每科分配时长 + 可解释理由"],
          ["调用条件", "多科目冲突、需要全局再分配时"],
          ["关键逻辑", "按「优先级 × 阶段紧迫度 × 连续失败严重度」计算每科权重，在每日可用时长内分配，带最低保底 0.5h"],
      ],
      widths=[3, 12])

h(doc, "3.2.2.4 Skill 4：interpret_feedback（反思理解）", 4)
table(doc,
      ["项", "内容"],
      [
          ["输入", "用户反思原文（自然语言）"],
          ["输出", "primary_cause / secondary_cause / weak_topics / emotion / evidence_summary"],
          ["调用条件", "① 用户提交新反思时；② diagnose 需要归因分析时"],
          ["关键逻辑", "用 LLM 从 8 个维度判定主因：方法不当/情绪干扰/精力不足/目标不清晰/时间投入不足/任务过载/基础薄弱/环境干扰；LLM 不可用时降级为关键词匹配"],
      ],
      widths=[3, 12])

h(doc, "3.2.2.5 Skill 5：proactive_scan（未来 7 天负荷扫描）", 4)
table(doc,
      ["项", "内容"],
      [
          ["输入", "未来 N 天计划 + 积压任务 backlog"],
          ["输出", "负荷预测 + 削峰填谷方案 + 降级决策"],
          ["调用条件", "计划生成后、需要事前负荷预测时"],
          ["关键逻辑", "负载上限 = 每日可用时长 × 1.2；超载日标红并给出削峰方案，把高峰任务平移至低峰日"],
      ],
      widths=[3, 12])

h(doc, "3.2.2.6 Skill 6：reschedule_for_calendar_change（课表变动重排）", 4)
table(doc,
      ["项", "内容"],
      [
          ["输入", "旧课表 + 新课表（version 变化即自动触发）"],
          ["输出", "冲突明细 + 重排后的新计划"],
          ["调用条件", "课表版本变化时（课表页「🔄 检测冲突并重排」）"],
          ["关键逻辑", "对比新旧课表、检测冲突、就近平移 / 跨天平移重排未来受影响任务，必要时触发 allocate 全局再分配"],
      ],
      widths=[3, 12])

h(doc, "3.2.2.7 Skill 7：decompose_goal（目标拆解）", 4)
table(doc,
      ["项", "内容"],
      [
          ["输入", "user_profile（goal / target_date）+ current_plan + learning_memory（+ rules）"],
          ["输出", "required_modules（需覆盖模块）/ covered_modules（已覆盖）/ missing_modules（缺失）/ coverage_rate / stages[]（阶段里程碑）/ current_stage / recommendation"],
          ["调用条件", "生成新计划前，检查「考研大目标」是否被完整拆解（对应命题背景「目标不清」）"],
          ["关键逻辑", "从目标关键词（数学一 / 数学二 / 408）识别需覆盖的知识模块，对照当前计划科目，指出「目标模块 vs 已覆盖模块」的缺口，并按下剩余天数拆「基础 / 强化 / 冲刺」三阶段里程碑"],
      ],
      widths=[3, 12])

h(doc, "3.2.2.8 Skill 8：aggregate_resources（资源聚合）", 4)
table(doc,
      ["项", "内容"],
      [
          ["输入", "弱知识点列表 weak_topics（来自 diagnose 或 memory）"],
          ["输出", "每个弱知识点对应的资源清单（视频 / 课后题 / 错题本 / 单词本）+ summary"],
          ["调用条件", "诊断出弱知识点后，把分散资料收拢到一条补齐路径（对应命题背景「资源分散」）"],
          ["关键逻辑", "按 RESOURCE_CATALOG 把弱知识点映射到具体资源（视频 / 课后题 / 错题本标签 / 单词本），未命中目录时走兜底资源，把散落的资料按弱项收拢"],
      ],
      widths=[3, 12])

h(doc, "3.2.3 技术创新点", 3)
bullet(doc, "补欠独立：重规划产生的「补欠」任务与常规任务分离，避免挤占正常进度；")
bullet(doc, "多级兜底：标红 → 拆分减时长 → 降级难度 → 换策略 + 三方向 → 人工确认；")
bullet(doc, "容错降级：LLM 不可用自动降级，保证任何环境可复现、可答辩；")
bullet(doc, "数据自洽：掌握度公式可回放（raw_data / formula / last_calculated），时长与时段自动校准。")

h(doc, "3.3 实施方案", 2)
h(doc, "3.3.1 实施路线图（已完成 / 答辩 / 后续）", 3)
bullet(doc, "已完成：memory.json 记忆、8 个 Skill、openJiuwen ReActAgent 调度、demo.py 逐日演示、HTML 前端、FastAPI 后端 9 接口、云端部署方案；")
bullet(doc, "答辩阶段：demo 演示录屏、解决方案文档、PPT、6 答题点对照表；")
bullet(doc, "后续：接入 Supabase 多用户、定时触发自动诊断、Web/小程序界面。")
h(doc, "3.3.2 demo.py 逐日演示说明（对应答题点 ③④）", 3)
p(doc, "场景：科目「数据结构」连续未完成，其余科目正常。逐日推进的完整调整链：")
table(doc,
      ["Day", "DS 打卡", "连续失败", "触发动作", "调整（before → after）"],
      [
          ["1", "✗ 未完成", "1", "无（仅记录）", "—"],
          ["2", "✗ 未完成", "2", "标红提醒（medium）", "不改变内容"],
          ["3", "✗ 未完成", "3", "diagnose → replan（第 1 次）", "减少时长 + 拆分任务"],
          ["4", "✗ 未完成", "4", "继续观察", "—"],
          ["5", "✗ 未完成", "5", "replan（第 2 次）", "降级内容难度"],
          ["6", "✗ 未完成", "6", "replan（第 3 次）", "换任务类型 + 给出三个方向（推荐 A）"],
      ],
      widths=[1.2, 2.8, 2.2, 4.6, 4.2])
p(doc, "关键案例（Day 3）计划生成前后对比：")
bullet(doc, "调整前：数据结构「单链表的插入与删除」，1.5h；")
bullet(doc, "调整后：数据结构「单链表的插入与删除（拆分：①看讲解视频 30min ②做 2 道基础题 30min）」，1.0h；")
bullet(doc, "理由：连续失败 3 天达到阈值，采取「减少时长 + 拆分任务」，降低启动门槛；")
bullet(doc, "依据（结构化）：连续失败 3 天 / 近 7 天完成率 0% / 反思原文「时间不够，写到一半就困了」「单链表插入删除还是卡，感觉是方法不对」/ 结论「方法不当」。")
h(doc, "3.3.3 资源配置方案", 3)
p(doc, "当前在本地运行，通过 cpolar 内网穿透供手机访问；生产环境规划部署到 Render + Supabase。")

h(doc, "3.4 需求匹配度分析", 2)
h(doc, "3.4.1 方案与命题对照（6 答题点逐条对应）", 3)
table(doc,
      ["答题点", "方案对应", "实现依据"],
      [
          ["① 记忆存储与更新", "memory.json + 打卡闭环", "update_memory_after_checkin 重算掌握度/连续失败/完成率"],
          ["② 至少 2 个 Skill", "8 个 Skill", "skills.py，输入/输出/调用条件见 3.2.2"],
          ["③ 动态调整案例", "连续 3 天未完成 → 减时长+拆分", "demo.py Day 3"],
          ["④ 计划前后对比", "before/after/reason/evidence", "adjustments[] 结构"],
          ["⑤ openJiuwen 作用", "AgentCard + ReActAgent + @tool", "agent.py 决策编排层"],
          ["⑥ 个性化与可解释性", "掌握度公式 + 证据链", "mastery_score 公式 + reason/evidence"],
      ],
      widths=[3.5, 5.5, 6])
h(doc, "3.4.2 匹配度评估", 3)
p(doc, "6 个答题点全部有明确的代码与演示支撑，无空泛描述；每个答题点均可现场演示 + 逐条解释 reason/evidence。")
h(doc, "3.4.3 可行性论证", 3)
p(doc, "技术可行（开源栈 + 确定性规则）、经济可行（离线可跑）、时间可行（已完成开发），详见 1.3.2。")

h(doc, "3.5 创新成效", 2)
bullet(doc, "形成了「打卡 → 记忆更新 → 诊断 → 重规划 → 新计划」的可演示闭环；")
bullet(doc, "每一次调整都可解释、可审计、可复现；")
bullet(doc, "具备「离线 / 无 key / 超时」全环境可跑通的鲁棒性。")
h(doc, "3.6 预期效益", 2)
bullet(doc, "对用户：摆脱手动重排，获得有依据的个性化动态计划；")
bullet(doc, "对企业（openJiuwen）：一个可复用的「规则 + 编排」领域落地范式；")
bullet(doc, "对团队：完整走通 Agent 应用从设计到落地的工程闭环。")
need(doc, "demo.py 运行截图 / 录屏", "6 个 Skill 前端结果卡片截图", "掌握度公式与 reason/evidence 的可视化截图")

# ============================================================
# 第四章 团队协作
# ============================================================
h(doc, "第四章　团队协作", 1)
h(doc, "4.1 团队介绍", 2)
p(doc, "【待补充：成员姓名、专业、年级、技术分工。】")
h(doc, "4.2 团队组织架构", 2)
p(doc, "【待补充：如「算法与 Agent 层 / 前端与产品层 / 文档与答辩层」的分工图。】")
h(doc, "4.3 团队与项目的关系", 2)
p(doc, "项目由团队成员完整自主开发，覆盖 Agent 决策层、规则引擎、前端交互、后端接口与部署，体现团队的技术栈覆盖与协作能力。")
h(doc, "4.4 团队与企业持续合作", 2)
p(doc, "【待补充：与命题企业 openJiuwen 相关的沟通、技术社区参与、后续合作意向等。】")
h(doc, "4.5 外部资源", 2)
p(doc, "【待补充：指导教师、实验室、算力 / 数据 / 场地等外部支持。】")
need(doc, "成员信息与分工", "组织架构图", "与企业沟通记录 / 合作意向", "外部资源说明")

# ============================================================
# 第五章 实施计划与里程碑
# ============================================================
h(doc, "第五章　实施计划与里程碑", 1)
h(doc, "5.1 总体实施规划", 2)
p(doc, "项目已按「需求分析 → 架构设计 → 规则引擎 → 前端后端 → 演示脚本 → 答辩材料」推进，当前进入答辩打磨阶段。")
h(doc, "5.2 关键里程碑（M1–M5）", 2)
table(doc,
      ["里程碑", "内容", "状态"],
      [
          ["M1", "记忆 Schema 与掌握度公式设计", "已完成"],
          ["M2", "8 个 Skill 规则引擎（skills.py）", "已完成"],
          ["M3", "openJiuwen ReActAgent 调度 + demo.py 逐日演示", "已完成"],
          ["M4", "HTML 前端 + FastAPI 后端（9 接口）", "已完成"],
          ["M5", "解决方案文档 / PPT / 录屏 / 答辩演练", "进行中"],
      ],
      widths=[2, 9, 4])
h(doc, "5.3 难点与重点", 2)
bullet(doc, "难点：如何保证「动态调整」可复现（用确定性规则解决）、如何保证「可解释」（用 reason/evidence 结构化解决）；")
bullet(doc, "重点：6 答题点全覆盖 + 现场可演示 + 可解释性讲清楚。")
h(doc, "5.4 资源配置计划", 2)
p(doc, "【待补充：人员投入、时间排期、所需算力 / 数据资源。】")
need(doc, "具体时间排期表", "里程碑完成的截图 / commit 记录佐证")

# ============================================================
# 第六章 风险分析与应对
# ============================================================
h(doc, "第六章　风险分析与应对", 1)
h(doc, "6.1 风险识别（技术 / 实施 / 合作 / 资源）", 2)
bullet(doc, "技术风险：LLM 依赖（已用降级机制化解）；框架 API 变动（锁定 openJiuwen 版本）；")
bullet(doc, "实施风险：答辩现场网络 / 演示环境不稳定（已保证离线可跑）；")
bullet(doc, "合作风险：与企业互动不足；")
bullet(doc, "资源风险：数据规模、算力预算有限。")
h(doc, "6.2 风险评估", 2)
table(doc,
      ["风险", "可能性", "影响", "等级"],
      [
          ["LLM 不可用", "中", "低（已降级）", "低"],
          ["现场网络 / 环境异常", "中", "中", "中"],
          ["框架 API 变更", "低", "中", "低"],
          ["数据规模不足", "高", "低（demo 足够）", "低"],
      ],
      widths=[6, 2.5, 3.5, 3])
h(doc, "6.3 风险应对措施", 2)
bullet(doc, "LLM 不可用：确定性规则 + 关键词降级，离线即可完整演示；")
bullet(doc, "现场异常：提前准备本地录屏 + 截图 + 离线 demo 脚本；")
bullet(doc, "框架变更：requirements.txt 锁定版本，README 记录兼容说明。")
need(doc, "更完整的风险清单（结合团队实际情况补充）")

# ============================================================
# 第七章 附件
# ============================================================
h(doc, "第七章　附件（佐证材料）", 1)
h(doc, "7.1 demo.py 运行输出（截图）", 2)
p(doc, "【待补充：demo.py 运行结果截图，重点截 Day 3 的 before/after 对比。】")
h(doc, "7.2 memory.json 数据结构说明", 2)
p(doc, "已具备，见 agent/docs/memory_schema.md：user_profile / current_plan / learning_memory（掌握度公式）/ rules / recent_records / replanning_log 六大部分。")
h(doc, "7.3 代码仓库链接", 2)
p(doc, "https://github.com/jeffreysgeorgy-crypto/kaoyan-checkin")
h(doc, "7.4 demo 演示录屏", 2)
p(doc, "【待补充：demo 演示录屏（建议 3–5 分钟，覆盖 8 个 Skill）。】")
h(doc, "7.5 6 个答题点对照表", 2)
p(doc, "已具备，见 1.3.1 / 3.4.1 对照表。")
h(doc, "7.6 HTML 打卡系统演示截图", 2)
p(doc, "【待补充：前端打卡 / 课表 / 错题本 / 番茄钟 / 统计 / 6 个 Skill 结果卡片截图。】")
need(doc, "demo.py 运行截图", "录屏文件", "前端各页面截图", "如需附 PPT 与答辩逐字稿一并整理")

doc.save(OUT)
print("✅ 已生成", OUT)
